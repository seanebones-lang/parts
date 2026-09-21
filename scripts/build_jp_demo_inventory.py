#!/usr/bin/env python3
"""Build JP-tailored synthetic demo inventory CSV from the official workbook.

Deterministic. No LLM. No external APIs.

Source (repo copy):
  data/jp_demo/JP_Transmission_Synthetic_Inventory.xlsx

Outputs:
  data/jp_demo/jp_demo_inventory_seed.csv
  data/jp_demo/jp_demo_inventory_seed.meta.json

Usage:
  PYTHONPATH=src /tmp/phase4-integrated-venv/bin/python scripts/build_jp_demo_inventory.py
  PYTHONPATH=src /tmp/phase4-integrated-venv/bin/python scripts/build_jp_demo_inventory.py --check
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import sys
from collections import Counter
from dataclasses import asdict, dataclass, field
from pathlib import Path

try:
    import openpyxl
except ImportError as exc:  # pragma: no cover
    raise SystemExit("openpyxl required: pip install openpyxl") from exc

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_XLSX = ROOT / "data" / "jp_demo" / "JP_Transmission_Synthetic_Inventory.xlsx"
DEFAULT_CSV = ROOT / "data" / "jp_demo" / "jp_demo_inventory_seed.csv"
DEFAULT_META = ROOT / "data" / "jp_demo" / "jp_demo_inventory_seed.meta.json"

# Importer contract
CSV_HEADERS = [
    "sku",
    "name",
    "transmission_family",
    "location",
    "qty",
    "transmission_variant",
    "category",
    "description",
    "oem_brand",
    "list_price",
    "condition",
    "bin",
    "part_type",
    "identifier_type",
    "identifier_value",
    "verification_status",
]

# Workbook warehouse label → existing DMS location codes (CHI-N / OHARE only).
WAREHOUSE_TO_CODE = {
    "main warehouse": "CHI-N",
    "converter cage": "CHI-N",
    "core yard": "CHI-N",
    "dismantling floor hold": "OHARE",
}

# Workbook condition → importer ALLOWED_CONDITIONS {new, used, rebuilt, core}
CONDITION_MAP = {
    "used / inspected": "used",
    "used / as-removed": "used",
    "used / good": "used",
    "hold for inspect": "used",
    "parts-only": "used",
    "rebuildable": "core",
    "as-removed": "core",
    "teardown candidate": "core",
}

# Free-form workbook part_type → resolver-known type when possible (else description only).
PART_TYPE_MAP = {
    "valve body": "valve body",
    "pump assembly": "pump",
    "input drum": "drum",
    "reverse input drum": "drum",
}

# Existing DEMO catalog rows we may top-up via preserve (exact name + family match required).
# Workbook JP-* SKUs stay as new pilot SKUs — we do not invent identity merges.
CANONICAL_TOPUP: list[dict[str, str | int]] = [
    {
        "sku": "4L60E-VB-01",
        "name": "4L60E Valve Body",
        "transmission_family": "4L60E",
        "transmission_variant": "4L60E",
        "part_type": "valve body",
        "location": "CHI-N",
        "qty": 12,
        "condition": "used",
        "bin": "JP-VB-A1",
        "description": "4L60E valve body · JP pilot top-up (canonical demo SKU)",
        "oem_brand": "GM",
        "category": "transmission_hard_parts",
        "verification_status": "unverified",
    },
    {
        "sku": "6R80-PUMP-01",
        "name": "6R80 Pump Assembly",
        "transmission_family": "6R80",
        "transmission_variant": "6R80",
        "part_type": "pump",
        "location": "CHI-N",
        "qty": 9,
        "condition": "used",
        "bin": "JP-PMP-A1",
        "description": "6R80 pump assembly · JP pilot top-up (canonical demo SKU)",
        "oem_brand": "Ford",
        "category": "transmission_hard_parts",
        "verification_status": "unverified",
    },
    {
        "sku": "6L80-PUMP-01",
        "name": "6L80 Transmission Pump Assembly",
        "transmission_family": "6L80",
        "transmission_variant": "6L80",
        "part_type": "pump",
        "location": "CHI-N",
        "qty": 6,
        "condition": "used",
        "bin": "JP-PMP-B1",
        "description": "OEM-style pump for 6L80/6L90 · JP pilot top-up",
        "oem_brand": "GM",
        "category": "transmission_hard_parts",
        "verification_status": "unverified",
    },
]


@dataclass
class BuildReport:
    source_path: str
    source_sha256: str
    workbook_inventory_rows: int = 0
    emitted_rows: int = 0
    unique_skus: int = 0
    skipped: list[dict[str, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    families: dict[str, int] = field(default_factory=dict)
    part_types: dict[str, int] = field(default_factory=dict)
    stock_states: dict[str, int] = field(default_factory=dict)
    locations: dict[str, int] = field(default_factory=dict)
    conditions: dict[str, int] = field(default_factory=dict)
    new_jp_skus: int = 0
    canonical_topup_rows: int = 0
    identifiers_with_casting: int = 0
    pricing_included: int = 0
    reserved_from_workbook_ignored: int = 0


def _sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _cell(row: tuple, idx: dict[str, int], key: str):
    i = idx.get(key)
    if i is None:
        return None
    return row[i]


def _map_condition(raw: str) -> str | None:
    key = (raw or "").strip().lower()
    return CONDITION_MAP.get(key)


def _map_location(raw: str) -> str | None:
    key = (raw or "").strip().lower()
    return WAREHOUSE_TO_CODE.get(key)


def _map_part_type(raw: str) -> str:
    key = (raw or "").strip().lower()
    return PART_TYPE_MAP.get(key, "")


def _stock_state(qty: int) -> str:
    if qty <= 0:
        return "out"
    if qty == 1:
        return "low"
    return "in"


def _ident_type(value: str) -> str:
    """Workbook column is casting_or_id — treat as casting provenance (not invented OEM)."""
    v = (value or "").strip()
    if not v:
        return ""
    # Ford-style part numbers still come from workbook casting_or_id column.
    return "casting"


def _description(part_name: str, part_type: str, notes: str, unit_code: str, apps: str) -> str:
    bits = [part_name or ""]
    pt = (part_type or "").strip()
    if pt and pt.lower() not in (part_name or "").lower():
        bits.append(pt)
    if unit_code and unit_code not in (part_name or ""):
        bits.append(f"unit {unit_code}")
    if notes:
        bits.append(str(notes).strip())
    # applications are descriptive context only — not fitment rows
    if apps:
        bits.append(f"typical apps: {str(apps).strip()[:160]}")
    return " · ".join(b for b in bits if b)


def build_rows(xlsx: Path) -> tuple[list[dict[str, str]], BuildReport]:
    if not xlsx.is_file():
        raise FileNotFoundError(f"workbook not found: {xlsx}")

    report = BuildReport(source_path=str(xlsx.resolve()), source_sha256=_sha256(xlsx))
    wb = openpyxl.load_workbook(xlsx, data_only=True, read_only=True)
    if "Inventory" not in wb.sheetnames:
        raise ValueError(f"Inventory sheet missing; sheets={wb.sheetnames}")
    ws = wb["Inventory"]
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter)
    headers = [str(h).strip() if h is not None else "" for h in header_row]
    idx = {h: i for i, h in enumerate(headers) if h}
    required = [
        "sku",
        "part_name",
        "family",
        "part_type",
        "qty_on_hand",
        "warehouse",
        "condition",
    ]
    missing = [c for c in required if c not in idx]
    if missing:
        raise ValueError(f"Inventory sheet missing columns: {missing}")

    out: list[dict[str, str]] = []
    seen_sku_loc: set[tuple[str, str]] = set()
    fam_c: Counter[str] = Counter()
    pt_c: Counter[str] = Counter()
    stock_c: Counter[str] = Counter()
    loc_c: Counter[str] = Counter()
    cond_c: Counter[str] = Counter()

    n_src = 0
    for row in rows_iter:
        if row is None or all(v is None or str(v).strip() == "" for v in row):
            continue
        n_src += 1
        sku = str(_cell(row, idx, "sku") or "").strip()
        status = str(_cell(row, idx, "status") or "Active").strip()
        name = str(_cell(row, idx, "part_name") or "").strip()
        family = str(_cell(row, idx, "family") or "").strip()
        unit_code = str(_cell(row, idx, "unit_code") or family).strip()
        part_type_raw = str(_cell(row, idx, "part_type") or "").strip()
        category = str(_cell(row, idx, "category") or "transmission_hard_parts").strip()
        casting = str(_cell(row, idx, "casting_or_id") or "").strip()
        cond_raw = str(_cell(row, idx, "condition") or "").strip()
        wh_raw = str(_cell(row, idx, "warehouse") or "").strip()
        bin_raw = str(_cell(row, idx, "bin") or "").strip()
        notes = str(_cell(row, idx, "notes") or "").strip()
        apps = str(_cell(row, idx, "typical_applications") or "").strip()
        make = str(_cell(row, idx, "make") or "").strip()
        price = _cell(row, idx, "unit_price")
        qty_raw = _cell(row, idx, "qty_on_hand")
        qty_res = _cell(row, idx, "qty_reserved")

        if not sku:
            report.skipped.append({"reason": "missing_sku", "row": str(n_src)})
            continue
        if status.lower() == "hold":
            # Keep Hold lots visible but do not exclude — they are valid pilot stock
            pass

        try:
            qty = int(qty_raw)
        except (TypeError, ValueError):
            report.skipped.append({"reason": "bad_qty", "sku": sku, "qty": repr(qty_raw)})
            continue
        if qty < 0:
            report.skipped.append({"reason": "negative_qty", "sku": sku, "qty": str(qty)})
            continue

        try:
            reserved_wb = int(qty_res or 0)
        except (TypeError, ValueError):
            reserved_wb = 0
        if reserved_wb:
            report.reserved_from_workbook_ignored += 1

        loc = _map_location(wh_raw)
        if not loc:
            report.skipped.append(
                {"reason": "unmapped_warehouse", "sku": sku, "warehouse": wh_raw}
            )
            continue

        cond = _map_condition(cond_raw)
        if not cond:
            report.skipped.append(
                {"reason": "unmapped_condition", "sku": sku, "condition": cond_raw}
            )
            continue

        key = (sku, loc)
        if key in seen_sku_loc:
            report.skipped.append(
                {"reason": "duplicate_sku_location", "sku": sku, "location": loc}
            )
            continue
        seen_sku_loc.add(key)

        pt_known = _map_part_type(part_type_raw)
        desc = _description(name, part_type_raw, notes, unit_code, apps)
        # Ensure known types appear in description for LIKE match
        if pt_known and pt_known not in desc.lower():
            desc = f"{desc} {pt_known}".strip()

        id_type = _ident_type(casting)
        id_val = casting if id_type else ""

        list_price = ""
        if price is not None and str(price).strip() != "":
            try:
                lp = float(price)
                if lp >= 0:
                    list_price = f"{lp:.2f}"
                    report.pricing_included += 1
            except (TypeError, ValueError):
                report.warnings.append(f"{sku}: non-numeric unit_price {price!r} omitted")

        rec = {
            "sku": sku,
            "name": name or sku,
            "transmission_family": family,
            "location": loc,
            "qty": str(qty),
            "transmission_variant": unit_code or family,
            "category": "transmission_hard_parts",
            "description": desc,
            "oem_brand": make,
            "list_price": list_price,
            "condition": cond,
            "bin": bin_raw,
            "part_type": pt_known,
            "identifier_type": id_type if id_val else "",
            "identifier_value": id_val,
            "verification_status": "unverified",
        }
        if id_val:
            report.identifiers_with_casting += 1

        out.append(rec)
        fam_c[family] += 1
        pt_c[part_type_raw or "(blank)"] += 1
        stock_c[_stock_state(qty)] += 1
        loc_c[loc] += 1
        cond_c[cond] += 1
        report.new_jp_skus += 1

    report.workbook_inventory_rows = n_src

    # Canonical top-ups (inventory on existing demo identity)
    for top in CANONICAL_TOPUP:
        sku = str(top["sku"])
        loc = str(top["location"])
        key = (sku, loc)
        if key in seen_sku_loc:
            report.warnings.append(f"canonical top-up skipped; collision {key}")
            continue
        seen_sku_loc.add(key)
        qty = int(top["qty"])
        rec = {
            "sku": sku,
            "name": str(top["name"]),
            "transmission_family": str(top["transmission_family"]),
            "location": loc,
            "qty": str(qty),
            "transmission_variant": str(top["transmission_variant"]),
            "category": str(top["category"]),
            "description": str(top["description"]),
            "oem_brand": str(top["oem_brand"]),
            "list_price": "",  # do not overwrite catalog list_price on preserve
            "condition": str(top["condition"]),
            "bin": str(top["bin"]),
            "part_type": str(top["part_type"]),
            "identifier_type": "",
            "identifier_value": "",
            "verification_status": "unverified",
        }
        out.append(rec)
        report.canonical_topup_rows += 1
        fam_c[str(top["transmission_family"])] += 1
        pt_c[str(top["part_type"])] += 1
        stock_c[_stock_state(qty)] += 1
        loc_c[loc] += 1
        cond_c[str(top["condition"])] += 1

    # Stable sort
    out.sort(key=lambda r: (r["transmission_family"], r["part_type"], r["sku"], r["location"]))

    report.emitted_rows = len(out)
    report.unique_skus = len({r["sku"] for r in out})
    report.families = dict(sorted(fam_c.items(), key=lambda kv: (-kv[1], kv[0])))
    report.part_types = dict(sorted(pt_c.items(), key=lambda kv: (-kv[1], kv[0])))
    report.stock_states = dict(stock_c)
    report.locations = dict(loc_c)
    report.conditions = dict(cond_c)
    wb.close()
    return out, report


def write_csv(rows: list[dict[str, str]], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_HEADERS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({h: r.get(h, "") for h in CSV_HEADERS})


def write_meta(report: BuildReport, path: Path, csv_path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = asdict(report)
    payload["csv_path"] = str(csv_path.resolve())
    payload["csv_sha256"] = _sha256(csv_path) if csv_path.is_file() else ""
    payload["notes"] = [
        "SYNTHETIC DEMO INVENTORY — not verified live JP stock.",
        "Workbook qty_reserved ignored; import leaves reserved_qty at 0.",
        "Fitment/interchange not generated from this CSV (canonical seed only).",
        "Identifiers come only from workbook casting_or_id column as casting.",
        "list_price from workbook unit_price is synthetic demo pricing.",
        "Locations mapped to CHI-N (Main Warehouse) and OHARE (Front Counter).",
    ]
    path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def validate_rows(rows: list[dict[str, str]]) -> list[str]:
    errs: list[str] = []
    seen: set[tuple[str, str]] = set()
    for i, r in enumerate(rows, start=2):
        sku = r.get("sku", "")
        loc = r.get("location", "")
        if not sku or not r.get("name") or not r.get("transmission_family") or not loc:
            errs.append(f"line {i}: missing required field")
        try:
            qty = int(r.get("qty", "0"))
        except ValueError:
            errs.append(f"line {i}: bad qty")
            continue
        if qty < 0:
            errs.append(f"line {i}: negative qty")
        if loc not in ("CHI-N", "OHARE"):
            errs.append(f"line {i}: invalid location {loc}")
        if r.get("condition") not in ("new", "used", "rebuilt", "core"):
            errs.append(f"line {i}: bad condition {r.get('condition')}")
        key = (sku, loc)
        if key in seen:
            errs.append(f"line {i}: duplicate {key}")
        seen.add(key)
        idt, idv = r.get("identifier_type", ""), r.get("identifier_value", "")
        if bool(idt) != bool(idv):
            errs.append(f"line {i}: identifier pair incomplete")
    return errs


def main(argv: list[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    ap.add_argument("--meta", type=Path, default=DEFAULT_META)
    ap.add_argument(
        "--check",
        action="store_true",
        help="Rebuild and assert identical output to existing CSV (byte-stable).",
    )
    args = ap.parse_args(argv)

    rows, report = build_rows(args.xlsx)
    errs = validate_rows(rows)
    if errs:
        print("VALIDATION ERRORS:", file=sys.stderr)
        for e in errs[:40]:
            print(" ", e, file=sys.stderr)
        return 2

    if args.check:
        if not args.csv.is_file():
            print(f"--check failed: missing {args.csv}", file=sys.stderr)
            return 3
        prev = args.csv.read_bytes()
        # write to temp buffer
        import io

        buf = io.StringIO()
        w = csv.DictWriter(buf, fieldnames=CSV_HEADERS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow({h: r.get(h, "") for h in CSV_HEADERS})
        new = buf.getvalue().encode("utf-8")
        if new != prev:
            print("--check FAILED: CSV output drifted", file=sys.stderr)
            return 4
        print("check ok: identical CSV", args.csv)
        print(
            json.dumps(
                {
                    "rows": report.emitted_rows,
                    "skus": report.unique_skus,
                    "skipped": len(report.skipped),
                }
            )
        )
        return 0

    write_csv(rows, args.csv)
    write_meta(report, args.meta, args.csv)
    print(
        json.dumps(
            {
                "csv": str(args.csv),
                "meta": str(args.meta),
                "rows": report.emitted_rows,
                "unique_skus": report.unique_skus,
                "skipped": len(report.skipped),
                "families": len(report.families),
                "source_sha256": report.source_sha256,
                "csv_sha256": _sha256(args.csv),
            },
            indent=2,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
